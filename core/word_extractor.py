"""Extracción de texto y tablas desde archivos Word (.docx)."""
from docx import Document


def extract_docx_text(file_path):
    doc = Document(file_path)
    return "\n".join(p.text for p in doc.paragraphs)


def extract_docx_tables(file_path):
    doc = Document(file_path)
    rows = []
    for table in doc.tables:
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
    return rows


def extract_docx_all(file_path):
    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)
